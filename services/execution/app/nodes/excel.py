import asyncio
import csv
import io
import logging
import re
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path
from typing import Any

import boto3
from openpyxl import load_workbook

from app.engine.context import RunContext
from app.engine.errors import NodeError

logger = logging.getLogger(__name__)


def normalise_header(raw: Any, position: int, taken: set[str]) -> str:
    text = "" if raw is None else str(raw).strip().lower()
    cleaned = re.sub(r"[^a-z0-9]+", "_", text).strip("_") or f"column_{position + 1}"
    if cleaned[0].isdigit():
        cleaned = f"col_{cleaned}"
    candidate, suffix = cleaned, 2
    while candidate in taken:
        candidate = f"{cleaned}_{suffix}"
        suffix += 1
    taken.add(candidate)
    return candidate


def build_headers(header_row) -> list[str]:
    taken: set[str] = set()
    return [normalise_header(cell, index, taken) for index, cell in enumerate(header_row)]


def to_json_value(value: Any) -> Any:
    if value is None or isinstance(value, (bool, int, float)):
        return value
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return float(value)
    return str(value)


def is_blank_row(row) -> bool:
    return all(cell is None or (isinstance(cell, str) and not cell.strip()) for cell in row)


def rows_to_items(rows, max_rows: int) -> list[dict]:
    iterator = iter(rows)
    header_row = next(iterator, None)
    if header_row is None:
        return []
    headers = build_headers(header_row)
    items: list[dict] = []
    for row in iterator:
        if is_blank_row(row):
            continue
        if len(items) >= max_rows:
            raise NodeError(f"File has more than {max_rows} rows")
        items.append({header: to_json_value(cell) for header, cell in zip(headers, row)})
    return items


def parse_xlsx(content: bytes, sheet: str | None, max_rows: int) -> list[dict]:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:
        raise NodeError("Could not open the Excel file") from None
    try:
        if sheet:
            if sheet not in workbook.sheetnames:
                raise NodeError(f"Sheet '{sheet}' was not found")
            worksheet = workbook[sheet]
        else:
            worksheet = workbook.worksheets[0]
        return rows_to_items(worksheet.iter_rows(values_only=True), max_rows)
    finally:
        workbook.close()


def parse_csv(content: bytes, max_rows: int) -> list[dict]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise NodeError("CSV file must be UTF-8 encoded") from None
    return rows_to_items(csv.reader(io.StringIO(text)), max_rows)


def is_safe_key(path: str) -> bool:
    return not path.startswith("/") and ".." not in Path(path).parts and "\\" not in path


def read_local_file(settings, path: str) -> bytes:
    if not settings.allow_local_files:
        raise NodeError("Local files are disabled on this server")
    base = Path(settings.local_files_dir).resolve()
    candidate = (base / path).resolve()
    if not candidate.is_relative_to(base) or not candidate.is_file():
        raise NodeError("File was not found in the local files folder")
    if candidate.stat().st_size > settings.max_file_bytes:
        raise NodeError("File is too large")
    return candidate.read_bytes()


def read_s3_object(settings, key: str) -> bytes:
    if not settings.s3_bucket:
        raise NodeError("S3 bucket is not configured on this server")
    try:
        client = boto3.client("s3", region_name=settings.aws_region)
        head = client.head_object(Bucket=settings.s3_bucket, Key=key)
        if head["ContentLength"] > settings.max_file_bytes:
            raise NodeError("File is too large")
        return client.get_object(Bucket=settings.s3_bucket, Key=key)["Body"].read(settings.max_file_bytes + 1)
    except NodeError:
        raise
    except Exception:
        logger.exception("S3 read failed for key %s", key)
        raise NodeError("Could not read the file from S3") from None


async def excel_read(node: dict, items: list[dict], ctx: RunContext) -> list[dict]:
    config = node["config"]
    path = config["path"].strip()
    if not is_safe_key(path):
        raise NodeError("File path is not allowed")
    lowered = path.lower()
    if not lowered.endswith((".xlsx", ".csv")):
        raise NodeError("Only .xlsx and .csv files are supported")

    reader = read_local_file if config["source"] == "local" else read_s3_object
    content = await asyncio.to_thread(reader, ctx.settings, path)
    max_rows = min(config.get("max_rows", 5000), ctx.settings.max_items_per_node)
    if lowered.endswith(".csv"):
        return await asyncio.to_thread(parse_csv, content, max_rows)
    return await asyncio.to_thread(parse_xlsx, content, config.get("sheet"), max_rows)
